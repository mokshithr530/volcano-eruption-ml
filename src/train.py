"""Train and evaluate eruption-status classifiers from the supplied Kilauea catalog."""
from __future__ import annotations

import json
import re
from datetime import date, datetime
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from sklearn.dummy import DummyClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.impute import SimpleImputer
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (accuracy_score, classification_report,
                             balanced_accuracy_score, cohen_kappa_score, confusion_matrix, f1_score,
                             precision_score, recall_score, roc_auc_score,
                             roc_curve)
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "raw"
OUT = ROOT / "outputs"
OUT.mkdir(exist_ok=True)


def duration_days(value) -> float:
    """Convert the chronology's human-readable durations into days."""
    s = str(value).strip().lower()
    if s in {"nan", "none", ""}:
        return 1.0
    m = re.search(r"([0-9]+(?:\.[0-9]+)?)", s)
    if not m:
        return 1.0
    n = float(m.group(1))
    if "year" in s or "yr" in s:
        return n * 365.25
    if "hour" in s or "hr" in s:
        return max(n / 24.0, 1 / 24.0)
    return max(n, 1 / 24.0)


def load_eruption_intervals(path: Path) -> list[tuple[pd.Timestamp, pd.Timestamp]]:
    """Read the Pu'u 'O'o chronology embedded in the workbook."""
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["Eruptions"]
    intervals = []
    for row in ws.iter_rows(values_only=True):
        event_id, start, duration = row[0], row[1], row[2]
        if not re.fullmatch(r"\d+[A-Za-z]?", str(event_id).strip()):
            continue
        # The worksheet also contains a summary table with text dates. The
        # detailed Pu'u 'O'o chronology stores true Excel date values.
        if not isinstance(start, (datetime, date)):
            continue
        start = pd.to_datetime(start, errors="coerce")
        if pd.isna(start):
            continue
        # Match the time span of the supplied earthquake catalog. The workbook
        # contains later chronology tables that are outside this experiment.
        if not (pd.Timestamp("1983-01-01") <= start <= pd.Timestamp("1986-12-31")):
            continue
        end = start + pd.Timedelta(days=duration_days(duration))
        intervals.append((start, end))
    if not intervals:
        raise ValueError("No eruption chronology rows were found")
    return intervals


def in_eruption(ts: pd.Timestamp, intervals) -> int:
    return int(any(start <= ts <= end for start, end in intervals))


def make_features(eq: pd.DataFrame, intervals) -> pd.DataFrame:
    eq = eq.sort_values("timestamp").reset_index(drop=True)
    times = eq["timestamp"].astype("int64").to_numpy()
    magnitudes = eq["magnitude"].to_numpy(float)
    rows = []
    windows = {"rate_1d": 1, "rate_7d": 7, "rate_30d": 30}
    for i, row in eq.iterrows():
        t = times[i]
        item = {
            "timestamp": row["timestamp"], "latitude": row["latitude"],
            "longitude": row["longitude"], "depth": row["depth"],
            "magnitude": row["magnitude"], "distance": row["distance"],
            "erupting": in_eruption(row["timestamp"], intervals),
        }
        for name, days in windows.items():
            left = t - int(days * 86400 * 1e9)
            prior = np.searchsorted(times[:i], left, side="left")
            item[name] = i - prior
            item[f"maxmag_{days}d"] = (magnitudes[prior:i].max()
                                        if i > prior else 0.0)
        rows.append(item)
    return pd.DataFrame(rows)


def choose_threshold(y_true, probabilities):
    """Choose a threshold on development data without looking at the test set."""
    candidates = np.linspace(0.05, 0.95, 181)
    scores = [(balanced_accuracy_score(y_true, probabilities >= t),
               f1_score(y_true, probabilities >= t, zero_division=0), t)
              for t in candidates]
    return max(scores)[2]


def evaluate(name, model, x_train, y_train, x_test, y_test, threshold=0.5):
    model.fit(x_train, y_train)
    prob = model.predict_proba(x_test)[:, 1] if hasattr(model, "predict_proba") else pred
    pred = (prob >= threshold).astype(int)
    metrics = {
        "model": name,
        "threshold": float(threshold),
        "accuracy": float(accuracy_score(y_test, pred)),
        "precision": float(precision_score(y_test, pred, zero_division=0)),
        "recall": float(recall_score(y_test, pred, zero_division=0)),
        "f1": float(f1_score(y_test, pred, zero_division=0)),
        "kappa": float(cohen_kappa_score(y_test, pred)),
        "roc_auc": float(roc_auc_score(y_test, prob)) if len(np.unique(y_test)) == 2 else None,
        "confusion_matrix": confusion_matrix(y_test, pred).tolist(),
        "classification_report": classification_report(y_test, pred, zero_division=0),
    }
    return model, metrics, prob


def main():
    eq = pd.read_csv(DATA / "puuoo_earthquakes.csv")
    eq = eq.rename(columns={"Date-time": "timestamp", "Latitude": "latitude",
                            "Longitude": "longitude", "Depth": "depth",
                            "Magnitude": "magnitude", "Distance": "distance"})
    # The source catalog uses US-style month/day/year timestamps.
    eq["timestamp"] = pd.to_datetime(eq["timestamp"], errors="coerce")
    eq = eq.dropna(subset=["timestamp", "magnitude", "depth", "latitude", "longitude"])
    intervals = load_eruption_intervals(DATA / "kilauea_eruptionHistory.xlsx")
    features = make_features(eq, intervals)
    features.to_csv(OUT / "engineered_features.csv", index=False)

    feature_cols = [c for c in features.columns if c not in {"timestamp", "erupting"}]
    cutoff_train = features["timestamp"].quantile(0.70)
    cutoff_dev = features["timestamp"].quantile(0.85)
    train = features[features.timestamp <= cutoff_train]
    dev = features[(features.timestamp > cutoff_train) & (features.timestamp <= cutoff_dev)]
    test = features[features.timestamp > cutoff_dev]
    x_train, y_train = train[feature_cols], train["erupting"]
    x_dev, y_dev = dev[feature_cols], dev["erupting"]
    x_test, y_test = test[feature_cols], test["erupting"]

    models = {
        "majority_baseline": DummyClassifier(strategy="most_frequent"),
        "logistic_regression": make_pipeline(SimpleImputer(), StandardScaler(),
                                               LogisticRegression(max_iter=2000, class_weight="balanced")),
        "random_forest": make_pipeline(SimpleImputer(), RandomForestClassifier(
            n_estimators=300, max_depth=15, min_samples_leaf=3,
            class_weight="balanced", random_state=42, n_jobs=-1)),
    }
    all_metrics, probabilities, fitted_models = [], {}, {}
    for name, model in models.items():
        model.fit(x_train, y_train)
        dev_prob = model.predict_proba(x_dev)[:, 1] if hasattr(model, "predict_proba") else model.predict(x_dev)
        threshold = choose_threshold(y_dev, dev_prob) if name != "majority_baseline" else 0.5
        _, metrics, prob = evaluate(name, model, x_train, y_train, x_test, y_test, threshold)
        all_metrics.append(metrics)
        probabilities[name] = prob
        fitted_models[name] = model
    rf = fitted_models["random_forest"].named_steps["randomforestclassifier"]
    importances = pd.DataFrame({"feature": feature_cols, "importance": rf.feature_importances_}).sort_values("importance", ascending=False)
    importances.to_csv(OUT / "feature_importance.csv", index=False)
    plt.figure(figsize=(7, 4))
    top = importances.head(8).sort_values("importance")
    plt.barh(top["feature"], top["importance"], color="#DD6C27")
    plt.xlabel("Mean decrease in impurity"); plt.title("Random Forest feature importance")
    plt.tight_layout(); plt.savefig(OUT / "feature_importance.png", dpi=180); plt.close()
    (OUT / "metrics.json").write_text(json.dumps({
        "dataset": {"rows": int(len(features)), "positive_rate": float(features.erupting.mean()),
                    "train_rows": int(len(train)), "test_rows": int(len(test)),
                    "dev_rows": int(len(dev)), "train_end": str(cutoff_train), "development_end": str(cutoff_dev)},
        "features": feature_cols, "models": all_metrics
    }, indent=2))

    plt.figure(figsize=(7, 5))
    for name, prob in probabilities.items():
        if len(np.unique(y_test)) == 2:
            fpr, tpr, _ = roc_curve(y_test, prob)
            auc = roc_auc_score(y_test, prob)
            plt.plot(fpr, tpr, label=f"{name.replace('_', ' ')} (AUC={auc:.2f})")
    plt.plot([0, 1], [0, 1], "k--", linewidth=1)
    plt.xlabel("False positive rate"); plt.ylabel("True positive rate")
    plt.title("Eruption-status classification on chronological test set")
    plt.legend(fontsize=8); plt.tight_layout(); plt.savefig(OUT / "roc_curves.png", dpi=180); plt.close()

    best = next(m for m in all_metrics if m["model"] == "random_forest")
    cm = np.array(best["confusion_matrix"])
    plt.figure(figsize=(4.5, 4))
    plt.imshow(cm, cmap="Reds")
    for (i, j), v in np.ndenumerate(cm): plt.text(j, i, str(v), ha="center", va="center")
    plt.xticks([0, 1], ["Repose", "Erupting"]); plt.yticks([0, 1], ["Repose", "Erupting"])
    plt.xlabel("Predicted"); plt.ylabel("Actual"); plt.title("Random Forest confusion matrix")
    plt.tight_layout(); plt.savefig(OUT / "confusion_matrix.png", dpi=180); plt.close()
    print(json.dumps({"rows": len(features), "positive_rate": float(features.erupting.mean()),
                      "models": [{k: v for k, v in m.items() if k not in {"classification_report", "confusion_matrix"}}
                                 for m in all_metrics]}, indent=2))


if __name__ == "__main__":
    main()
